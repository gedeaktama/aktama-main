# -*- coding: utf-8 -*-
import docx
import re

doc = docx.Document(r'D:\OneDrive - Universitas Parnaraya\GEDE ERIK AKTAMA\TOT\Outline Buku Ajar KKA\Buku Kelas 5 SD\Buku_Panduan_Guru_Koding_dan_KA_Kelas_5_SD.docx')

print(f"Buku Guru total paragraphs: {len(doc.paragraphs)}")

bullets = []
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if '•' in text:
        clean = re.sub(r'^[•\s]+', '', text).strip()
        bullets.append((i, clean))

print(f"Buku Guru bullets: {len(bullets)}")
for i, c in bullets[:30]:
    print(f"  p[{i:3d}] {c[:80]}")
