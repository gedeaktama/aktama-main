# -*- coding: utf-8 -*-
"""
Professional Publishing Compiler for Buku Panduan Guru SD Kelas 5
Standard: Permendikdasmen No. 13 Tahun 2025, Pusbuk BSKAP, UNESCO B5 Format
"""
import os, sys, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import guru_data_part1
import guru_bab_1_2_3
import guru_bab_4_5_6

tot_dir = r"D:\OneDrive - Universitas Parnaraya\GEDE ERIK AKTAMA\TOT\Outline Buku Ajar KKA\Buku Kelas 5 SD"
local_dir = r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-5-sd"

# Combine full guru markdown
full_guru_md = "\n\n".join([
    guru_data_part1.GURU_FRONT_MATTER,
    guru_bab_1_2_3.GURU_BAB_1,
    guru_bab_1_2_3.GURU_BAB_2,
    guru_bab_1_2_3.GURU_BAB_3,
    guru_bab_4_5_6.GURU_BAB_4,
    guru_bab_4_5_6.GURU_BAB_5,
    guru_bab_4_5_6.GURU_BAB_6
])

# Save markdown to both locations
for base in [tot_dir, local_dir]:
    with open(os.path.join(base, "Buku_Panduan_Guru_Koding_dan_KA_Kelas_5_SD.md"), "w", encoding="utf-8") as f:
        f.write(full_guru_md.strip() + "\n")
print("Professional Markdown Buku Panduan Guru written.")

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def apply_guru_callout_box(paragraph, bg_hex="FEF3C7", border_hex="D97706"):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = parse_xml(f'''
        <w:pBdr {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="14" w:color="{border_hex}"/>
        </w:pBdr>
    ''')
    pPr.append(pBdr)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    pPr.append(shd)

def build_pro_guru_docx(md_content, output_name):
    doc = docx.Document()
    
    # 1. Page Dimension: UNESCO B5 (15.5 cm x 23.0 cm)
    for s in doc.sections:
        s.page_width = Pt(439.45)
        s.page_height = Pt(652.05)
        s.top_margin = Pt(56.7)
        s.bottom_margin = Pt(56.7)
        s.left_margin = Pt(56.7)
        s.right_margin = Pt(56.7)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Buku Panduan Guru: Koding & Kecerdasan Artifisial — Kelas 5 SD (Fase C)")
        frun.font.size = Pt(8.5)
        frun.font.name = 'Arial'
        frun.font.color.rgb = RGBColor(146, 64, 14)

    # 2. Formal Title Page
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(36)
    p_t.paragraph_format.space_after = Pt(8)
    r_t = p_t.add_run("BUKU PANDUAN GURU")
    r_t.font.name = 'Arial'
    r_t.font.bold = True
    r_t.font.size = Pt(20)
    r_t.font.color.rgb = RGBColor(180, 83, 9)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_sub = p_sub.add_run("KODING DAN KECERDASAN ARTIFISIAL\nUNTUK SD/MI KELAS V (FASE C)")
    r_sub.font.name = 'Arial'
    r_sub.font.bold = True
    r_sub.font.size = Pt(13.5)
    r_sub.font.color.rgb = RGBColor(30, 41, 59)

    p_pill = doc.add_paragraph()
    p_pill.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pill.paragraph_format.space_after = Pt(40)
    r_pill = p_pill.add_run("PEDOMAN OPERASIONAL PEMBELAJARAN MENDALAM (DEEP LEARNING)\nBERKESADARAN (MINDFUL) • BERMAKNA (MEANINGFUL) • MENGGEMBIRAKAN (JOYFUL)")
    r_pill.font.name = 'Arial'
    r_pill.font.italic = True
    r_pill.font.size = Pt(9)
    r_pill.font.color.rgb = RGBColor(71, 85, 105)

    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_auth.paragraph_format.space_after = Pt(70)
    r_auth = p_auth.add_run("Penulis:\nGede Erik Aktama, M.Kom\n(Fasilitator Nasional KKA Kemendikdasmen RI)")
    r_auth.font.name = 'Arial'
    r_auth.font.bold = True
    r_auth.font.size = Pt(11)
    r_auth.font.color.rgb = RGBColor(15, 23, 42)

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("PUSAT PERBUKUAN\nBADAN STANDAR, KURIKULUM, DAN ASESMEN PENDIDIKAN\nKEMENTERIAN PENDIDIKAN DASAR DAN MENENGAH RI\n2025 / 2026")
    r_inst.font.name = 'Arial'
    r_inst.font.size = Pt(9)
    r_inst.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    lines = md_content.split("\n")
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        rows_data = []
        for r in table_rows:
            parts = [c.strip() for c in r.split("|")[1:-1]]
            if parts and not all(set(c).issubset({'-', ':', ' '}) for c in parts):
                rows_data.append(parts)
        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.style = 'Table Grid'
            for ri, row in enumerate(rows_data):
                for ci, cell_text in enumerate(row):
                    if ci < num_cols:
                        cell = tbl.cell(ri, ci)
                        cell.text = cell_text
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.space_before = Pt(2)
                            for r in p.runs:
                                r.font.name = 'Arial'
                                r.font.size = Pt(8.5)
                        if ri == 0:
                            set_cell_background(cell, "FEF3C7")
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.bold = True
                                    r.font.color.rgb = RGBColor(146, 64, 14)
        in_table = False
        table_rows = []

    for line in lines:
        sline = line.strip()

        if sline.startswith("|") and sline.endswith("|"):
            in_table = True
            table_rows.append(sline)
            continue
        else:
            if in_table:
                flush_table()

        if not sline:
            continue

        if sline.startswith("# "):
            if not sline.startswith("# PANDUAN GURU:"):
                doc.add_page_break()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(sline[2:])
            run.font.name = 'Arial'
            run.font.bold = True
            run.font.size = Pt(14.5)
            run.font.color.rgb = RGBColor(146, 64, 14)
        elif sline.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(sline[3:])
            run.font.name = 'Arial'
            run.font.bold = True
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(180, 83, 9)
        elif sline.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline[4:])
            run.font.name = 'Arial'
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif sline.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(sline[5:])
            run.font.name = 'Arial'
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif sline.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(sline.lstrip("> "))
            run.font.name = 'Arial'
            run.font.size = Pt(9)
            run.font.italic = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            apply_guru_callout_box(p, bg_hex="FFFBEB", border_hex="F59E0B")
        elif sline.startswith("* ") or sline.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline[2:])
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)
        elif any(sline.startswith(f"{i}. ") for i in range(1, 40)):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline.split(". ", 1)[1] if ". " in sline else sline)
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)
        elif sline.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 38)
            run.font.name = 'Arial'
            run.font.color.rgb = RGBColor(226, 232, 240)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(sline)
            run.font.name = 'Arial'
            run.font.size = Pt(9.5)

    if in_table:
        flush_table()

    for base in [tot_dir, local_dir]:
        out_f = os.path.join(base, output_name)
        doc.save(out_f)
    print(f"Professional Guru DOCX Saved: {output_name}")

build_pro_guru_docx(full_guru_md, "Buku_Panduan_Guru_Koding_dan_KA_Kelas_5_SD.docx")
print("GURU PUBLISHING BUILD COMPLETED SUCCESSFULLY!")
