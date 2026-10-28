# -*- coding: utf-8 -*-
"""
Compiler Buku Siswa Koding & KA Kelas 5 SD
Format: UNESCO B5 (15.5 x 23 cm)
Target: Markdown (.md) and Microsoft Word (.docx)
"""
import os, sys, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import sd5_data

tot_dir = r"D:\OneDrive - Universitas Parnaraya\GEDE ERIK AKTAMA\TOT\Outline Buku Ajar KKA\Buku Kelas 5 SD"
workspace_dir = r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-5-sd"
cover_img_path = os.path.join(tot_dir, "assets", "cover_buku_siswa_sd5.jpg")

os.makedirs(os.path.join(tot_dir, "assets"), exist_ok=True)
os.makedirs(os.path.join(workspace_dir, "assets"), exist_ok=True)

# Combine markdown
full_md = "\n\n".join([
    sd5_data.FRONT_MATTER,
    sd5_data.BAB_1,
    sd5_data.BAB_2,
    sd5_data.BAB_3,
    sd5_data.BAB_4,
    sd5_data.BAB_5,
    sd5_data.BAB_6
])

# Save Markdown
for base in [tot_dir, workspace_dir]:
    md_path = os.path.join(base, "Buku_Siswa_Koding_dan_KA_Kelas_5_SD.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(full_md.strip() + "\n")
print("Markdown Buku Siswa saved to both locations.")

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_callout_box(paragraph, bg_hex="F1F5F9", border_color="0284C7"):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = parse_xml(f'''
        <w:pBdr {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="12" w:color="{border_color}"/>
        </w:pBdr>
    ''')
    pPr.append(pBdr)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    pPr.append(shd)

def create_docx(md_content, output_name, cover_path):
    doc = docx.Document()
    
    # Page setup UNESCO B5 (15.5 x 23 cm)
    for section in doc.sections:
        section.page_width = Pt(439.45)
        section.page_height = Pt(652.05)
        section.top_margin = Pt(56.7)
        section.bottom_margin = Pt(56.7)
        section.left_margin = Pt(56.7)
        section.right_margin = Pt(56.7)
        
        # Header and Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Koding & Kecerdasan Artifisial — Kelas 5 SD (Fase C)")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    # Add Cover Page
    if cover_path and os.path.exists(cover_path):
        p_cov = doc.add_paragraph()
        p_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cov.paragraph_format.space_before = Pt(0)
        p_cov.paragraph_format.space_after = Pt(0)
        run = p_cov.add_run()
        run.add_picture(cover_path, width=Inches(4.9))
        doc.add_page_break()

    lines = md_content.split("\n")
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        rows_data = []
        for r in table_rows:
            parts = [c.strip() for c in r.split("|")[1:-1]]
            if parts and not all(set(c).issubset({'-', ':', ' '}) for c in parts):
                rows_data.append(parts)
        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.style = 'Table Grid'
            for ri, row in enumerate(rows_data):
                for ci, cell_text in enumerate(row):
                    if ci < num_cols:
                        cell = tbl.cell(ri, ci)
                        cell.text = cell_text
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.space_before = Pt(2)
                            for r in p.runs:
                                r.font.size = Pt(8.5)
                        if ri == 0:
                            set_cell_background(cell, "E2E8F0")
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.bold = True
        in_table = False
        table_rows = []

    for line in lines:
        sline = line.strip()
        if sline.startswith("|") and sline.endswith("|"):
            in_table = True
            table_rows.append(sline)
            continue
        else:
            if in_table:
                flush_table()

        if not sline:
            continue

        if sline.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(sline[2:])
            run.font.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif sline.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(sline[3:])
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(30, 58, 138)
        elif sline.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline[4:])
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(13, 148, 136)
        elif sline.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(sline[5:])
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif sline.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(sline.lstrip("> "))
            run.font.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 41, 59)
            set_callout_box(p, bg_hex="F8FAFC", border_color="0EA5E9")
        elif sline.startswith("* ") or sline.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline[2:])
            run.font.size = Pt(9.5)
        elif any(sline.startswith(f"{i}. ") for i in range(1, 40)):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline.split(". ", 1)[1] if ". " in sline else sline)
            run.font.size = Pt(9.5)
        elif sline.startswith("```"):
            pass
        elif sline.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 38)
            run.font.color.rgb = RGBColor(203, 213, 225)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(sline)
            run.font.size = Pt(9.5)

    if in_table:
        flush_table()

    for base in [tot_dir, workspace_dir]:
        out_path = os.path.join(base, output_name)
        doc.save(out_path)
    print(f"DOCX Buku Siswa saved: {output_name}")

create_docx(full_md, "Buku_Siswa_Koding_dan_KA_Kelas_5_SD.docx", cover_img_path)
print("COMPILATION FINISHED SUCCESSFULLY!")
