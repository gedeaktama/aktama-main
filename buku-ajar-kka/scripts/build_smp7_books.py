# -*- coding: utf-8 -*-
"""
Master Builder for Official UNESCO Template Books - Grade 7 SMP (Fase D)
Produces:
1. Buku_Siswa_Koding_dan_KA_Kelas_7_SMP.docx & .md
2. Buku_Panduan_Guru_Koding_dan_KA_Kelas_7_SMP.docx & .md
Base Template: 'TEMPLATE BUKU AJAR UKURAN UNESCO.docx' (15,5 cm x 23,0 cm)
"""
import os, sys, re, shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

sys.path.insert(0, r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\scripts")
from docx_advanced_parser import add_styled_runs, clean_latex_symbols
from smp7_naskah_frontmatter import FRONT_MATTER_SMP7, BAGIAN_AKHIR_SMP7
from smp7_naskah_bab1 import BAB_1_SMP7
from smp7_naskah_bab2 import BAB_2_SMP7
from smp7_naskah_bab3 import BAB_3_SMP7
from smp7_naskah_bab4 import BAB_4_SMP7
from smp7_guru_panduan import full_guru_smp7_md

template_src = r"D:\OneDrive - Universitas Parnaraya\GEDE ERIK AKTAMA\TOT\Outline Buku Ajar KKA\template buku\TEMPLATE BUKU AJAR UKURAN UNESCO.docx"
tot_dir = r"D:\OneDrive - Universitas Parnaraya\GEDE ERIK AKTAMA\TOT\Outline Buku Ajar KKA\Buku Kelas 7 SMP"
local_dir = r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-7-smp"
assets_local = os.path.join(local_dir, "assets")
assets_tot = os.path.join(tot_dir, "assets")

os.makedirs(tot_dir, exist_ok=True)
os.makedirs(local_dir, exist_ok=True)
os.makedirs(assets_local, exist_ok=True)
os.makedirs(assets_tot, exist_ok=True)

cover_img_path = os.path.join(assets_local, "cover_buku_siswa_smp7.jpg")

full_siswa_smp7_md = "\n\n".join([
    FRONT_MATTER_SMP7,
    BAB_1_SMP7,
    BAB_2_SMP7,
    BAB_3_SMP7,
    BAB_4_SMP7,
    BAGIAN_AKHIR_SMP7
])

# Save Markdown files to both locations
print("Saving Markdown files...")
for base in [tot_dir, local_dir]:
    siswa_md_path = os.path.join(base, "Buku_Siswa_Koding_dan_KA_Kelas_7_SMP.md")
    with open(siswa_md_path, "w", encoding="utf-8") as f:
        f.write(full_siswa_smp7_md.strip() + "\n")
    print(f"Saved: {siswa_md_path}")

    guru_md_path = os.path.join(base, "Buku_Panduan_Guru_Koding_dan_KA_Kelas_7_SMP.md")
    with open(guru_md_path, "w", encoding="utf-8") as f:
        f.write(full_guru_smp7_md.strip() + "\n")
    print(f"Saved: {guru_md_path}")

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def apply_callout_box(paragraph, bg_hex, border_hex):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = parse_xml(f'''
        <w:pBdr {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="14" w:color="{border_hex}"/>
        </w:pBdr>
    ''')
    pPr.append(pBdr)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    pPr.append(shd)

def prepare_template_doc(template_path, new_footer_text):
    doc = docx.Document(template_path)
    for section in doc.sections:
        footer = section.footer
        if footer.paragraphs:
            p_foot = footer.paragraphs[0]
            p_foot.text = ""
            p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_foot = p_foot.add_run(f"{new_footer_text}  |  ")
            r_foot.font.name = 'Times New Roman'
            r_foot.font.size = Pt(8.5)
            r_foot.font.italic = True
            r_foot.font.color.rgb = RGBColor(100, 116, 139)
            fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
            p_foot._p.append(fldSimple)

    sectPr = doc.sections[-1]._sectPr
    for child in list(doc._body._element):
        if child != sectPr:
            doc._body._element.remove(child)

    return doc

def build_full_book(md_content, output_filename, is_guru=False, cover_path=None):
    print(f"\nBuilding DOCX: {output_filename} (Guru={is_guru})...")
    footer_title = "Buku Panduan Guru: KKA Kelas VII SMP/MTs (Fase D)" if is_guru else "Koding & Kecerdasan Artifisial — Kelas VII SMP/MTs (Fase D)"
    doc = prepare_template_doc(template_src, footer_title)

    # Cover for Student Book
    if cover_path and os.path.exists(cover_path) and not is_guru:
        p_cov = doc.add_paragraph()
        p_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cov.paragraph_format.space_before = Pt(0)
        p_cov.paragraph_format.space_after = Pt(0)
        r_cov = p_cov.add_run()
        r_cov.add_picture(cover_path, width=Inches(4.9))
        doc.add_page_break()

    # Front Title Page for Teacher Book
    if is_guru:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(36)
        p_t.paragraph_format.space_after = Pt(8)
        r_t = p_t.add_run("BUKU PANDUAN GURU")
        r_t.font.name = 'Times New Roman'; r_t.font.bold = True; r_t.font.size = Pt(20)
        r_t.font.color.rgb = RGBColor(146, 64, 14)

        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(20)
        r_sub = p_sub.add_run("KODING DAN KECERDASAN ARTIFISIAL\nUNTUK SMP/MTs KELAS VII (FASE D)")
        r_sub.font.name = 'Times New Roman'; r_sub.font.bold = True; r_sub.font.size = Pt(13.5)
        r_sub.font.color.rgb = RGBColor(30, 41, 59)

        p_pill = doc.add_paragraph()
        p_pill.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_pill.paragraph_format.space_after = Pt(40)
        r_pill = p_pill.add_run("PEDOMAN OPERASIONAL PEMBELAJARAN MENDALAM (DEEP LEARNING)\nBERKESADARAN (MINDFUL) • BERMAKNA (MEANINGFUL) • MENGGEMBIRAKAN (JOYFUL)")
        r_pill.font.name = 'Times New Roman'; r_pill.font.italic = True; r_pill.font.size = Pt(9)
        r_pill.font.color.rgb = RGBColor(71, 85, 105)

        p_auth = doc.add_paragraph()
        p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_auth.paragraph_format.space_after = Pt(70)
        r_auth = p_auth.add_run("Penulis:\nGede Erik Aktama, S.Pd., M.Pd.\n(Fasilitator Nasional KKA Kemendikdasmen RI)\n(Dosen Lektor & Kaprodi Sistem Komputer, Universitas Parna Raya)")
        r_auth.font.name = 'Times New Roman'; r_auth.font.bold = True; r_auth.font.size = Pt(11)
        r_auth.font.color.rgb = RGBColor(15, 23, 42)

        p_inst = doc.add_paragraph()
        p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_inst = p_inst.add_run("PUSAT PERBUKUAN\nBADAN STANDAR, KURIKULUM, DAN ASESMEN PENDIDIKAN\nKEMENTERIAN PENDIDIKAN DASAR DAN MENENGAH RI\n2025 / 2026")
        r_inst.font.name = 'Times New Roman'; r_inst.font.size = Pt(9)
        r_inst.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_page_break()

    lines = md_content.split("\n")
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        rows_data = []
        for r in table_rows:
            parts = [c.strip() for c in r.split("|")[1:-1]]
            if parts and not all(set(c).issubset({'-', ':', ' '}) for c in parts):
                rows_data.append(parts)
        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.style = 'Table Grid'
            for ri, row in enumerate(rows_data):
                for ci, cell_text in enumerate(row):
                    if ci < num_cols:
                        cell = tbl.cell(ri, ci)
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_after = Pt(2)
                        p.paragraph_format.space_before = Pt(2)
                        if ri == 0:
                            set_cell_background(cell, "FEF3C7" if is_guru else "E0F2FE")
                            add_styled_runs(p, cell_text, base_font_size=8.5,
                                            base_color=RGBColor(146, 64, 14) if is_guru else RGBColor(3, 105, 161),
                                            font_name='Times New Roman')
                            for r in p.runs:
                                r.bold = True
                        else:
                            add_styled_runs(p, cell_text, base_font_size=8.5,
                                            base_color=RGBColor(30, 41, 59), font_name='Times New Roman')
        in_table = False
        table_rows = []

    current_section = "NORMAL"

    for line in lines:
        sline = line.strip()

        if sline.startswith("|") and sline.endswith("|"):
            in_table = True
            table_rows.append(sline)
            continue
        else:
            if in_table:
                flush_table()

        if not sline:
            continue

        # Embedded Images [[IMAGE: filename.png | Caption]]
        if sline.startswith("[[IMAGE:") and "]]" in sline:
            tag_content = sline.replace("[[IMAGE:", "").replace("]]", "").strip()
            parts = [p.strip() for p in tag_content.split("|")]
            img_filename = parts[0]
            caption = parts[1] if len(parts) > 1 else ""
            img_full = os.path.join(assets_local, img_filename)
            if not os.path.exists(img_full):
                img_full = os.path.join(assets_tot, img_filename)
            if os.path.exists(img_full):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(3)
                r_img = p_img.add_run()
                img_width = Inches(1.5) if "foto_penulis" in img_filename else Inches(4.6)
                r_img.add_picture(img_full, width=img_width)
                if caption:
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(10)
                    r_cap = p_cap.add_run(caption)
                    r_cap.font.name = 'Times New Roman'
                    r_cap.font.size = Pt(8.5)
                    r_cap.font.italic = True
                    r_cap.font.color.rgb = RGBColor(71, 85, 105)
            continue

        clean_heading = sline.replace("**", "").replace("*", "")
        clean_upper = clean_heading.upper()

        # Section Context Detection
        if clean_heading.startswith("# ") or clean_heading.startswith("## ") or clean_heading.startswith("### ") or clean_heading.startswith("#### "):
            if "DAFTAR ISI" in clean_upper:
                current_section = "DAFTAR_ISI"
            elif "DAFTAR GAMBAR" in clean_upper or "DAFTAR TABEL" in clean_upper:
                current_section = "DAFTAR_GAMBAR"
            elif "IDENTITAS PENERBITAN" in clean_upper:
                current_section = "IDENTITAS"
            elif "INDEKS KATA" in clean_upper:
                current_section = "INDEKS"
            elif any(clean_upper.startswith(x) for x in ["# BAB", "## BAB", "### BAB", "# SEMESTER", "## SEMESTER", "# BAGIAN", "## BAGIAN", "## KATA PENGANTAR", "## PRAKATA", "## PETUNJUK", "## DAFTAR TOKOH", "## MENGENAL TOKOH", "## PETA KONSEP", "### AYO", "### LKPD", "### PROFIL PENULIS", "### DAFTAR PUSTAKA", "## A.", "## B.", "## C.", "## D.", "## E.", "## F.", "## G.", "## H.", "## I."]):
                current_section = "NORMAL"

            # Render Headings
            if clean_heading.startswith("# "):
                if not (clean_heading.startswith("# KODING") or clean_heading.startswith("# BUKU PANDUAN GURU")):
                    doc.add_page_break()
                p = doc.add_paragraph()
                if "BAB " in clean_heading:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(clean_heading[2:])
                run.font.name = 'Times New Roman'; run.font.bold = True; run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(15, 23, 42)
            elif clean_heading.startswith("## "):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run(clean_heading[3:])
                run.font.name = 'Times New Roman'; run.font.bold = True; run.font.size = Pt(12)
                run.font.color.rgb = RGBColor(180, 83, 9) if is_guru else RGBColor(2, 132, 199)
            elif clean_heading.startswith("### "):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(clean_heading[4:])
                run.font.name = 'Times New Roman'; run.font.bold = True; run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(30, 41, 59) if is_guru else RGBColor(13, 148, 136)
            elif clean_heading.startswith("#### "):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                run = p.add_run(clean_heading[5:])
                run.font.name = 'Times New Roman'; run.font.bold = True; run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(71, 85, 105)
            continue

        # Section 1: DAFTAR ISI (Clean Indented, Zero Bullets)
        if current_section == "DAFTAR_ISI":
            if sline.startswith("---"):
                current_section = "NORMAL"
                continue
            clean_item = re.sub(r'^[\*\-\s]+', '', sline).strip()
            if not clean_item:
                continue
            if clean_item.startswith("SEMESTER") or clean_item.startswith("BAGIAN AKHIR") or clean_item.startswith("Bagian "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(clean_item)
                run.font.name = 'Times New Roman'; run.font.bold = True; run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(2, 132, 199)
            elif clean_item.startswith("BAB ") or clean_item.startswith("Bab "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(1.5)
                run = p.add_run(clean_item)
                run.font.name = 'Times New Roman'; run.font.bold = True; run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(15, 23, 42)
            elif re.match(r'^[A-I]\.\s', clean_item) or clean_item.startswith("Misi "):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.38)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(clean_item)
                run.font.name = 'Times New Roman'; run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(51, 65, 85)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.18)
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                run = p.add_run(clean_item)
                run.font.name = 'Times New Roman'; run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(51, 65, 85)
            continue

        # Section 2: DAFTAR GAMBAR & TABEL
        if current_section == "DAFTAR_GAMBAR":
            if sline.startswith("---"):
                current_section = "NORMAL"
                continue
            clean_item = re.sub(r'^[\*\-\s]+', '', sline).strip()
            if not clean_item:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = -Inches(0.28)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            m = re.match(r'^(Gambar\s+[\d\.]+:|Tabel\s+[\d\.]+:)(.*)', clean_item)
            if m:
                r1 = p.add_run(m.group(1))
                r1.font.name = 'Times New Roman'; r1.font.bold = True; r1.font.size = Pt(9.5)
                r1.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(2, 132, 199)
                r2 = p.add_run(m.group(2))
                r2.font.name = 'Times New Roman'; r2.font.size = Pt(9.5)
                r2.font.color.rgb = RGBColor(51, 65, 85)
            else:
                add_styled_runs(p, clean_item, base_font_size=9.5, base_color=RGBColor(51, 65, 85), font_name='Times New Roman')
            continue

        # Section 3: IDENTITAS PENERBITAN
        if current_section == "IDENTITAS":
            if sline.startswith("---"):
                current_section = "NORMAL"
                continue
            clean_item = re.sub(r'^[\*\-\s]+', '', sline).strip()
            if not clean_item:
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(2)
            add_styled_runs(p, clean_item, base_font_size=9.5, base_color=RGBColor(30, 41, 59), font_name='Times New Roman')
            continue

        # Section 4: INDEKS KATA
        if current_section == "INDEKS":
            if sline.startswith("---"):
                current_section = "NORMAL"
                continue
            clean_item = re.sub(r'^[\*\-\s]+', '', sline).strip()
            if not clean_item:
                continue
            letter = clean_item.replace("#", "").replace("*", "").strip()
            if len(letter) == 1 and letter.isalpha():
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.left_indent = Inches(0)
                run = p.add_run(letter)
                run.font.name = 'Times New Roman'; run.font.bold = True; run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(2, 132, 199)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                add_styled_runs(p, clean_item, base_font_size=9, base_color=RGBColor(51, 65, 85), font_name='Times New Roman')
            continue

        # Normal Body Processing

        # Callout Boxes (> quote)
        if sline.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            raw_text = sline.lstrip("> ")
            add_styled_runs(p, raw_text, base_font_size=9.5, base_color=RGBColor(30, 41, 59), font_name='Times New Roman')
            for r in p.runs:
                r.italic = True
            if is_guru:
                apply_callout_box(p, bg_hex="FFFBEB", border_hex="F59E0B")
            else:
                apply_callout_box(p, bg_hex="F0F9FF", border_hex="0284C7")
            continue

        # Checkboxes (* [ ] or [ ]) -> Clean Ballot Box (□)
        if sline.startswith("* [ ]") or sline.startswith("- [ ]") or sline.startswith("[ ]"):
            clean_text = re.sub(r'^[\*\-\s]*\[\s*\]\s*', '', sline).strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = -Inches(0.18)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r_box = p.add_run("□  ")
            r_box.font.name = 'Times New Roman'; r_box.font.size = Pt(11); r_box.font.bold = True
            r_box.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(2, 132, 199)
            add_styled_runs(p, clean_text, base_font_size=10, base_color=RGBColor(30, 41, 59), font_name='Times New Roman')
            continue

        # Fix bullet then number/letter
        m_bnum = re.match(r'^[\*\-]\s+([0-9]+|[A-Za-z])[\.\)]\s+(.*)', sline)
        if m_bnum:
            prefix = m_bnum.group(1)
            rest = m_bnum.group(2)
            sline = f"{prefix}. {rest}"

        # Multiple-choice exam options: A. , B. , C. , D.
        m_opt = re.match(r'^([A-D])[\.\)]\s+(.*)', sline)
        if m_opt:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = -Inches(0.2)
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.15
            r_lbl = p.add_run(f"{m_opt.group(1)}.  ")
            r_lbl.font.name = 'Times New Roman'; r_lbl.font.size = Pt(10); r_lbl.font.bold = True
            r_lbl.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(2, 132, 199)
            add_styled_runs(p, m_opt.group(2), base_font_size=10, base_color=RGBColor(30, 41, 59), font_name='Times New Roman')
            continue

        # Numbered list: 1. 2. 10.
        m_num = re.match(r'^([0-9]+)[\.\)]\s+(.*)', sline)
        if m_num:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = -Inches(0.2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r_num = p.add_run(f"{m_num.group(1)}.  ")
            r_num.font.name = 'Times New Roman'; r_num.font.size = Pt(10); r_num.font.bold = True
            r_num.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(2, 132, 199)
            add_styled_runs(p, m_num.group(2), base_font_size=10, base_color=RGBColor(30, 41, 59), font_name='Times New Roman')
            continue

        # Bulleted list (* or -)
        if sline.startswith("* ") or sline.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = -Inches(0.15)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r_bullet = p.add_run("•  ")
            r_bullet.font.name = 'Times New Roman'; r_bullet.font.size = Pt(10); r_bullet.font.bold = True
            r_bullet.font.color.rgb = RGBColor(146, 64, 14) if is_guru else RGBColor(2, 132, 199)
            add_styled_runs(p, sline[2:], base_font_size=10, base_color=RGBColor(30, 41, 59), font_name='Times New Roman')
            continue

        # Divider (---)
        if sline.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_div = p.add_run("•  •  •")
            r_div.font.name = 'Times New Roman'; r_div.font.size = Pt(10)
            r_div.font.color.rgb = RGBColor(203, 213, 225)
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2.5)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_styled_runs(p, sline, base_font_size=10, base_color=RGBColor(30, 41, 59), font_name='Times New Roman')

    if in_table:
        flush_table()

    # Save DOCX to both locations
    local_out = os.path.join(local_dir, output_filename)
    doc.save(local_out)
    print(f"Saved local DOCX: {local_out}")

    tot_out = os.path.join(tot_dir, output_filename)
    try:
        shutil.copy2(local_out, tot_out)
        print(f"Copied DOCX to OneDrive: {tot_out}")
    except Exception as e:
        print(f"Warning copying {output_filename} to OneDrive: {e}")

if __name__ == "__main__":
    print("Compiling Grade 7 SMP Books...")
    # 1. Compile Buku Siswa
    build_full_book(full_siswa_smp7_md, "Buku_Siswa_Koding_dan_KA_Kelas_7_SMP.docx", is_guru=False, cover_path=cover_img_path)

    # 2. Compile Buku Guru
    build_full_book(full_guru_smp7_md, "Buku_Panduan_Guru_Koding_dan_KA_Kelas_7_SMP.docx", is_guru=True)

    print("\nAll Grade 7 SMP Books successfully compiled!")
