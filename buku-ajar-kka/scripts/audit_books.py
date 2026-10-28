# -*- coding: utf-8 -*-
"""
Audit Profesional Buku Siswa & Guru KKA Kelas 5 SD
Memeriksa: format, konsistensi, kelengkapan, font, gambar, tabel, dll.
"""
import docx, re, os

def audit_book(doc_path, book_name):
    doc = docx.Document(doc_path)
    print(f'\n{"="*70}')
    print(f'  AUDIT PROFESIONAL: {book_name}')
    print(f'{"="*70}')
    print(f'File: {doc_path}')
    print(f'Ukuran File: {os.path.getsize(doc_path):,} bytes')
    print(f'Total Paragraf: {len(doc.paragraphs)}')
    print(f'Total Tabel: {len(doc.tables)}')

    # Count images
    xml_content = doc._part.blob.decode('utf-8', errors='ignore')
    drawing_count = len(re.findall(r'<w:drawing', xml_content))
    print(f'Total Gambar Tertanam: {drawing_count}')

    # Page dimensions
    sec = doc.sections[0]
    w_cm = sec.page_width.inches * 2.54
    h_cm = sec.page_height.inches * 2.54
    print(f'Ukuran Halaman: {w_cm:.1f} x {h_cm:.1f} cm (Target: 15.5 x 23.0)')

    # Margins
    left_cm = sec.left_margin.inches * 2.54
    right_cm = sec.right_margin.inches * 2.54
    top_cm = sec.top_margin.inches * 2.54
    bottom_cm = sec.bottom_margin.inches * 2.54
    print(f'Margin: L={left_cm:.1f} R={right_cm:.1f} T={top_cm:.1f} B={bottom_cm:.1f} cm')

    # Footer/Header
    for i, section in enumerate(doc.sections):
        if section.footer.paragraphs:
            ft = section.footer.paragraphs[0].text
            if ft.strip():
                print(f'Footer (Section {i}): "{ft[:100]}"')
        if section.header.paragraphs:
            ht = section.header.paragraphs[0].text
            if ht.strip():
                print(f'Header (Section {i}): "{ht[:100]}"')

    # Font audit
    fonts_used = {}
    font_sizes = {}
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.name:
                fonts_used[r.font.name] = fonts_used.get(r.font.name, 0) + 1
            if r.font.size:
                sz = r.font.size.pt
                font_sizes[sz] = font_sizes.get(sz, 0) + 1
    print(f'\nFont yang Digunakan:')
    for fn, cnt in sorted(fonts_used.items(), key=lambda x: -x[1]):
        print(f'  {fn}: {cnt} runs')
    print(f'Ukuran Font:')
    for fs, cnt in sorted(font_sizes.items()):
        print(f'  {fs} pt: {cnt} runs')

    # Issue detection
    issues = {
        'double_star': [],
        'image_tag': [],
        'code_fence': [],
        'raw_hash': [],
        'raw_bullet_star': [],
        'empty_runs': 0,
        'orphan_numbers': [],
        'duplicate_headings': [],
    }

    all_headings = []
    prev_text = ''
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        
        if '**' in t:
            issues['double_star'].append((i, t[:120]))
        if '[[IMAGE' in t:
            issues['image_tag'].append((i, t[:120]))
        if t.startswith('```'):
            issues['code_fence'].append((i, t[:120]))
        if t.startswith('# ') or t.startswith('## ') or t.startswith('### '):
            issues['raw_hash'].append((i, t[:120]))
        if t.startswith('* **') or t.startswith('- **'):
            issues['raw_bullet_star'].append((i, t[:120]))
        
        # Detect runs with no text
        for r in p.runs:
            if r.text.strip() == '' and r.font.size:
                issues['empty_runs'] += 1
        
        # Detect duplicate headings
        if t and t == prev_text and len(t) > 10:
            issues['duplicate_headings'].append((i, t[:120]))
        prev_text = t
        
        # Collect headings for structure check
        if p.runs and p.runs[0].bold and len(t) > 5 and not t.startswith('•'):
            all_headings.append(t)

    print(f'\n--- MASALAH FORMAT ---')
    print(f'  Tanda ** mentah: {len(issues["double_star"])}')
    for line, txt in issues['double_star'][:5]:
        print(f'    Para {line}: {txt}')
    
    print(f'  Tag [[IMAGE tidak terproses: {len(issues["image_tag"])}')
    for line, txt in issues['image_tag'][:5]:
        print(f'    Para {line}: {txt}')
    
    print(f'  Blok kode ``` mentah: {len(issues["code_fence"])}')
    for line, txt in issues['code_fence'][:5]:
        print(f'    Para {line}: {txt}')
    
    print(f'  Heading # mentah: {len(issues["raw_hash"])}')
    for line, txt in issues['raw_hash'][:5]:
        print(f'    Para {line}: {txt}')
    
    print(f'  Heading duplikat: {len(issues["duplicate_headings"])}')
    for line, txt in issues['duplicate_headings'][:5]:
        print(f'    Para {line}: {txt}')

    # Table audit
    print(f'\n--- TABEL ---')
    for idx, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)
        header = ' | '.join(c.text[:25] for c in table.rows[0].cells) if rows > 0 else 'KOSONG'
        print(f'  Tabel {idx+1}: {rows} baris x {cols} kolom | Header: {header[:100]}')
        # Check empty cells
        empty_cells = 0
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() == '':
                    empty_cells += 1
        if empty_cells > 0:
            print(f'    !! {empty_cells} sel kosong terdeteksi')

    # Content structure check
    print(f'\n--- STRUKTUR KONTEN ---')
    bab_found = [h for h in all_headings if 'BAB ' in h.upper()]
    print(f'  Bab Terdeteksi: {len(bab_found)}')
    for b in bab_found:
        print(f'    - {b}')

    # Daftar Pustaka check
    pustaka_found = any('DAFTAR PUSTAKA' in p.text.upper() for p in doc.paragraphs)
    print(f'  Daftar Pustaka: {"ADA" if pustaka_found else "TIDAK ADA !!!"}')
    
    glosarium_found = any('GLOSARIUM' in p.text.upper() for p in doc.paragraphs)
    print(f'  Glosarium: {"ADA" if glosarium_found else "TIDAK ADA !!!"}')
    
    profil_found = any('PROFIL PENULIS' in p.text.upper() for p in doc.paragraphs)
    print(f'  Profil Penulis: {"ADA" if profil_found else "TIDAK ADA !!!"}')
    
    indeks_found = any('INDEKS KATA' in p.text.upper() for p in doc.paragraphs)
    print(f'  Indeks Kata: {"ADA" if indeks_found else "TIDAK ADA !!!"}')

    evaluasi_found = any('EVALUASI KOMPREHENSIF' in p.text.upper() or 'SUMATIF AKHIR' in p.text.upper() for p in doc.paragraphs)
    print(f'  Evaluasi Sumatif Akhir Tahun: {"ADA" if evaluasi_found else "TIDAK ADA !!!"}')

    # Line spacing audit
    spacing_issues = 0
    for p in doc.paragraphs:
        pf = p.paragraph_format
        if pf.line_spacing and pf.line_spacing < 1.0:
            spacing_issues += 1
    print(f'\n  Paragraf dgn spasi < 1.0: {spacing_issues}')

    # Empty paragraph clusters (>3 in a row)
    empty_streak = 0
    max_streak = 0
    for p in doc.paragraphs:
        if p.text.strip() == '':
            empty_streak += 1
            max_streak = max(max_streak, empty_streak)
        else:
            empty_streak = 0
    print(f'  Cluster paragraf kosong terpanjang: {max_streak}')

    # Word count estimate
    word_count = sum(len(p.text.split()) for p in doc.paragraphs)
    print(f'\n  Estimasi Jumlah Kata: {word_count:,}')
    print(f'  Estimasi Jumlah Halaman (250 kata/hal): ~{word_count//250} halaman')

    return issues

# Audit Buku Siswa
siswa_path = r'd:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-5-sd\Buku_Siswa_Koding_dan_KA_Kelas_5_SD.docx'
siswa_issues = audit_book(siswa_path, 'BUKU SISWA KKA KELAS 5 SD')

# Audit Buku Guru
guru_path = r'd:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-5-sd\Buku_Panduan_Guru_Koding_dan_KA_Kelas_5_SD.docx'
guru_issues = audit_book(guru_path, 'BUKU PANDUAN GURU KKA KELAS 5 SD')

print(f'\n{"="*70}')
print('  RINGKASAN TEMUAN AUDIT')
print(f'{"="*70}')
total_issues_siswa = (len(siswa_issues['double_star']) + len(siswa_issues['image_tag']) + 
                      len(siswa_issues['code_fence']) + len(siswa_issues['raw_hash']) +
                      len(siswa_issues['duplicate_headings']))
total_issues_guru = (len(guru_issues['double_star']) + len(guru_issues['image_tag']) + 
                     len(guru_issues['code_fence']) + len(guru_issues['raw_hash']) +
                     len(guru_issues['duplicate_headings']))
print(f'Buku Siswa - Total masalah format: {total_issues_siswa}')
print(f'Buku Guru  - Total masalah format: {total_issues_guru}')
