# -*- coding: utf-8 -*-
"""
Compiler Buku Panduan Guru Koding & KA Kelas 5 SD
Format: UNESCO B5 (15.5 x 23 cm)
Target: Markdown (.md) and Microsoft Word (.docx)
"""
import os, sys, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

import guru_data_part1
import guru_bab_1_2_3
import guru_bab_4_5_6

tot_dir = r"D:\OneDrive - Universitas Parnaraya\GEDE ERIK AKTAMA\TOT\Outline Buku Ajar KKA\Buku Kelas 5 SD"
workspace_dir = r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-5-sd"

os.makedirs(tot_dir, exist_ok=True)
os.makedirs(workspace_dir, exist_ok=True)

# Combine markdown
full_guru_md = "\n\n".join([
    guru_data_part1.GURU_FRONT_MATTER,
    guru_bab_1_2_3.GURU_BAB_1,
    guru_bab_1_2_3.GURU_BAB_2,
    guru_bab_1_2_3.GURU_BAB_3,
    guru_bab_4_5_6.GURU_BAB_4,
    guru_bab_4_5_6.GURU_BAB_5,
    guru_bab_4_5_6.GURU_BAB_6
])

# Save Markdown to both locations
for base in [tot_dir, workspace_dir]:
    md_path = os.path.join(base, "Buku_Panduan_Guru_Koding_dan_KA_Kelas_5_SD.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(full_guru_md.strip() + "\n")
print("Markdown Buku Panduan Guru saved.")

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_callout_box(paragraph, bg_hex="FEF3C7", border_color="D97706"):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = parse_xml(f'''
        <w:pBdr {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="12" w:color="{border_color}"/>
        </w:pBdr>
    ''')
    pPr.append(pBdr)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_hex}"/>')
    pPr.append(shd)

def create_guru_docx(md_content, output_name):
    doc = docx.Document()
    
    # Page setup UNESCO B5 (15.5 x 23 cm)
    for section in doc.sections:
        section.page_width = Pt(439.45)
        section.page_height = Pt(652.05)
        section.top_margin = Pt(56.7)
        section.bottom_margin = Pt(56.7)
        section.left_margin = Pt(56.7)
        section.right_margin = Pt(56.7)
        
        # Header and Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Buku Panduan Guru: KKA Kelas 5 SD (Fase C) — Kemendikdasmen")
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 53, 15)

    # Title Page
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(8)
    run_t = p_title.add_run("BUKU PANDUAN GURU")
    run_t.font.bold = True
    run_t.font.size = Pt(20)
    run_t.font.color.rgb = RGBColor(180, 83, 9)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    run_s = p_sub.add_run("KODING DAN KECERDASAN ARTIFISIAL\nUNTUK SD/MI KELAS V (FASE C)")
    run_s.font.bold = True
    run_s.font.size = Pt(14)
    run_s.font.color.rgb = RGBColor(30, 41, 59)

    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_after = Pt(40)
    run_b = p_badge.add_run("PEDOMAN OPERASIONAL PEMBELAJARAN MENDALAM (DEEP LEARNING)\nBERKESADARAN • BERMAKNA • MENGGEMBIRAKAN")
    run_b.font.italic = True
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(71, 85, 105)

    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_after = Pt(80)
    run_a = p_author.add_run("Penulis:\nGede Erik Aktama, M.Kom\n(Fasilitator Nasional KKA Kemendikdasmen RI)")
    run_a.font.bold = True
    run_a.font.size = Pt(11)
    run_a.font.color.rgb = RGBColor(15, 23, 42)

    p_pub = doc.add_paragraph()
    p_pub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_p = p_pub.add_run("PUSAT PERBUKUAN\nBADAN STANDAR, KURIKULUM, DAN ASESMEN PENDIDIKAN\nKEMENTERIAN PENDIDIKAN DASAR DAN MENENGAH RI\n2025 / 2026")
    run_p.font.size = Pt(9)
    run_p.font.color.rgb = RGBColor(100, 116, 139)
    
    doc.add_page_break()

    lines = md_content.split("\n")
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        rows_data = []
        for r in table_rows:
            parts = [c.strip() for c in r.split("|")[1:-1]]
            if parts and not all(set(c).issubset({'-', ':', ' '}) for c in parts):
                rows_data.append(parts)
        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.style = 'Table Grid'
            for ri, row in enumerate(rows_data):
                for ci, cell_text in enumerate(row):
                    if ci < num_cols:
                        cell = tbl.cell(ri, ci)
                        cell.text = cell_text
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.space_before = Pt(2)
                            for r in p.runs:
                                r.font.size = Pt(8.5)
                        if ri == 0:
                            set_cell_background(cell, "FEF3C7")
                            for p in cell.paragraphs:
                                for r in p.runs:
                                    r.font.bold = True
        in_table = False
        table_rows = []

    for line in lines:
        sline = line.strip()
        if sline.startswith("|") and sline.endswith("|"):
            in_table = True
            table_rows.append(sline)
            continue
        else:
            if in_table:
                flush_table()

        if not sline:
            continue

        if sline.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(sline[2:])
            run.font.bold = True
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(120, 53, 15)
        elif sline.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(sline[3:])
            run.font.bold = True
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(180, 83, 9)
        elif sline.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(sline[4:])
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif sline.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(sline[5:])
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif sline.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(sline.lstrip("> "))
            run.font.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 41, 59)
            set_callout_box(p, bg_hex="FFFBEB", border_color="F59E0B")
        elif sline.startswith("* ") or sline.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline[2:])
            run.font.size = Pt(9.5)
        elif any(sline.startswith(f"{i}. ") for i in range(1, 40)):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(sline.split(". ", 1)[1] if ". " in sline else sline)
            run.font.size = Pt(9.5)
        elif sline.startswith("```"):
            pass
        elif sline.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 38)
            run.font.color.rgb = RGBColor(226, 232, 240)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(sline)
            run.font.size = Pt(9.5)

    if in_table:
        flush_table()

    for base in [tot_dir, workspace_dir]:
        out_path = os.path.join(base, output_name)
        doc.save(out_path)
    print(f"DOCX Buku Guru saved: {output_name}")

create_guru_docx(full_guru_md, "Buku_Panduan_Guru_Koding_dan_KA_Kelas_5_SD.docx")
print("GURU COMPILATION FINISHED SUCCESSFULLY!")
