# -*- coding: utf-8 -*-
"""
Script to write full content for Buku Siswa Kelas 5 SD
"""
import os, sys, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

tot_dir = r"D:\OneDrive - Universitas Parnaraya\GEDE ERIK AKTAMA\TOT\Outline Buku Ajar KKA\Buku Kelas 5 SD"
workspace_dir = r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-5-sd"
cover_img_path = os.path.join(tot_dir, "assets", "cover_buku_siswa_sd5.jpg")

os.makedirs(os.path.join(tot_dir, "assets"), exist_ok=True)
os.makedirs(os.path.join(workspace_dir, "assets"), exist_ok=True)

# Helper function to save file to both locations
def save_md(filename, content):
    for base in [tot_dir, workspace_dir]:
        full_path = os.path.join(base, filename)
        with open(full_path, "w", encoding="utf-8") as f:
            f.write(content.strip() + "\n")
    print(f"MD Saved: {filename}")

def build_docx_from_markdown(md_text, docx_filename, cover_image=None):
    doc = docx.Document()
    
    # Set UNESCO B5 Page size: 15.5 cm x 23 cm (approx 439.45 pt x 652.05 pt)
    for section in doc.sections:
        section.page_width = Pt(439.45)
        section.page_height = Pt(652.05)
        section.top_margin = Pt(56.7)
        section.bottom_margin = Pt(56.7)
        section.left_margin = Pt(56.7)
        section.right_margin = Pt(56.7)

    # Add Cover Image if exists
    if cover_image and os.path.exists(cover_image):
        p_cov = doc.add_paragraph()
        p_cov.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p_cov.add_run()
        run.add_picture(cover_image, width=Inches(4.8))
        doc.add_page_break()

    # Parse markdown lines into docx
    lines = md_text.split("\n")
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        rows_data = []
        for r in table_rows:
            parts = [c.strip() for c in r.split("|")[1:-1]]
            if parts and not all(set(c).issubset({'-', ':', ' '}) for c in parts):
                rows_data.append(parts)
        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.style = 'Table Grid'
            for ri, row in enumerate(rows_data):
                for ci, cell_text in enumerate(row):
                    if ci < num_cols:
                        cell = tbl.cell(ri, ci)
                        cell.text = cell_text
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(2)
                            p.paragraph_format.space_before = Pt(2)
        in_table = False
        table_rows = []

    for line in lines:
        sline = line.strip()
        if sline.startswith("|") and sline.endswith("|"):
            in_table = True
            table_rows.append(sline)
            continue
        else:
            if in_table:
                flush_table()

        if not sline:
            continue

        if sline.startswith("# "):
            p = doc.add_paragraph(sline[2:], style='Heading 1')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif sline.startswith("## "):
            doc.add_paragraph(sline[3:], style='Heading 1')
        elif sline.startswith("### "):
            doc.add_paragraph(sline[4:], style='Heading 2')
        elif sline.startswith("#### "):
            doc.add_paragraph(sline[5:], style='Heading 3')
        elif sline.startswith("* ") or sline.startswith("- "):
            doc.add_paragraph(sline[2:], style='List Paragraph')
        elif any(sline.startswith(f"{i}. ") for i in range(1, 30)):
            doc.add_paragraph(sline, style='List Paragraph')
        elif sline.startswith(">"):
            p = doc.add_paragraph(sline.lstrip("> "), style='Normal')
            p.paragraph_format.left_indent = Inches(0.25)
        elif sline.startswith("---"):
            pass
        else:
            doc.add_paragraph(sline, style='Normal')

    if in_table:
        flush_table()

    for base in [tot_dir, workspace_dir]:
        full_docx = os.path.join(base, docx_filename)
        doc.save(full_docx)
    print(f"DOCX Saved: {docx_filename}")

# Part 1: Front Matter and Chapter 1
front_matter = """# KODING DAN KECERDASAN ARTIFISIAL
## BUKU TEKS PELAJARAN SISWA SD/MI KELAS V
### Mata Pelajaran Pilihan — Fase C

**Penulis:**  
Gede Erik Aktama, M.Kom  
*(Fasilitator Nasional Koding dan Kecerdasan Artifisial Kemendikdasmen)*  

**Penelaah & Editor:**  
Tim Pengembang Kurikulum KKA Dikdasmen  

**Penerbit:**  
Pusat Perbukuan  
Badan Standar, Kurikulum, dan Asesmen Pendidikan  
Kementerian Pendidikan Dasar dan Menengah Republik Indonesia  
Tahun 2025 / 2026  

---

## HAK CIPTA & KETENTUAN PENERBITAN
Judul Buku: Koding dan Kecerdasan Artifisial untuk SD/MI Kelas V  
Penulis: Gede Erik Aktama, M.Kom  
Penyunting: Tim Fasilitator KKA  
Desain Sampul & Ilustrasi: Tim Desain Visual Kreatif  
Tata Letak: Format Standar UNESCO (15,5 cm x 23 cm)  
Ukuran Buku: 15,5 cm x 23 cm (B5 UNESCO)  
Font Isi: Arial / Calibri 11 pt, Spasi 1.15  

*Diterbitkan oleh Kementerian Pendidikan Dasar dan Menengah Republik Indonesia sesuai dengan ketentuan UU No. 3 Tahun 2017 tentang Sistem Perbukuan dan Permendikdasmen No. 13 Tahun 2025 tentang Kurikulum Nasional.*

---

## KATA PENGANTAR
Puji syukur kita panjatkan ke hadirat Tuhan Yang Maha Esa, karena atas rahmat dan karunia-Nya, buku teks pelajaran **Koding dan Kecerdasan Artifisial untuk SD/MI Kelas V** ini dapat hadir di hadapan para peserta didik dan pendidik di seluruh penjuru Indonesia.

Di era digital dan kecerdasan artifisial saat ini, keterampilan abad ke-21 seperti berpikir komputasional, pemecahan masalah sistematis, kreativitas, dan literasi teknologi menjadi bekal esensial bagi generasi penerus bangsa. Sejalan dengan visi Asta Cita pemerintah dalam memperkuat pembangunan sumber daya manusia unggul, Kementerian Pendidikan Dasar dan Menengah menghadirkan mata pelajaran Koding dan Kecerdasan Artifisial (KKA) sebagai mata pelajaran pilihan mulai dari jenjang Sekolah Dasar pada Fase C (Kelas V dan VI).

Buku ini disusun dengan pendekatan **Pembelajaran Mendalam (*Deep Learning Approach*)** yang menekankan suasana belajar yang **berkesadaran (*mindful*)**, **bermakna (*meaningful*)**, dan **menggembirakan (*joyful*)**. Pembelajaran dirancang secara seimbang melalui aktivitas tanpa komputer (*unplugged*) dan aktivitas dengan perangkat digital (*plugged*). Nilai-nilai Profil Pelajar Pancasila dan prinsip keadaban digital diletakkan sebagai fondasi utama agar peserta didik tidak hanya cakap secara teknis, tetapi juga memiliki empati dan budi pekerti luhur.

Kami menyampaikan apresiasi dan terima kasih yang mendalam kepada penulis, tim penelaah, fasilitator, dan seluruh pihak yang telah mencurahkan dedikasi terbaiknya. Semoga buku ini dapat menjadi sahabat setia anak-anak Indonesia dalam menjelajahi dunia sains komputer yang menakjubkan.

Jakarta, 2025  
**Kepala Pusat Perbukuan Kemendikdasmen**

---

## PRAKATA PENULIS
Halo anak-anak Indonesia yang hebat dan cerdas!

Selamat datang di petualangan baru yang sangat seru dan mengasyikkan! Di Kelas V ini, kalian akan memulai perjalanan istimewa bersama buku **Koding dan Kecerdasan Artifisial**. 

Mungkin kalian bertanya-tanya, *"Apa itu koding? Apakah kecerdasan artifisial itu seperti robot di film kartun? Apakah belajar koding itu sulit?"* Jawabannya adalah: belajar koding itu sangat menyenangkan! Belajar koding sama seperti bermain menyusun balok puzzle dan melatih otak kita untuk berpikir runtut, rapi, dan kreatif. 

Di dalam buku ini, kalian tidak akan belajar sendirian. Kalian akan ditemani oleh sahabat-sahabat baru yang ceria: **Erik** yang gemar berpetualang dan memecahkan teka-teki, **Siti** yang tekun dan suka menggambar digital, **Made** yang suka bercerita, dan **Kibo**, si robot pintar yang ramah dan siap membantu kalian memahami rahasia teknologi.

Kalian akan belajar cara memecahkan masalah sehari-hari, menjaga rahasia data pribadi di internet, melatih komputer mengenali gambar buah dan hewan, hingga membuat game dan cerita animasi interaktif kalian sendiri menggunakan Scratch!

Ingatlah pesan penting ini: Komputer dan robot cerdas diciptakan oleh manusia. Komputer hanyalah alat bantu. Yang paling istimewa dan paling pintar adalah **akal budi, hati nurani, dan kasih sayang yang ada di dalam diri kalian sebagai manusia**. Gunakanlah teknologi ini untuk berbuat kebaikan, membantu orang tua, menolong teman, dan menjaga kelestarian alam Indonesia.

Selamat belajar, berpetualang, dan berkreasi!

Salam hangat,  
**Gede Erik Aktama, M.Kom**  
*(Fasilitator Nasional KKA Kemendikdasmen)*

---

## PETUNJUK PENGGUNAAN BUKU SISWA
Agar petualangan belajarmu berjalan lancar dan menggembirakan, perhatikan ikon-ikon panduan belajar berikut:

1. **Ayo Menyimak (Apersepsi):** Cerita pembuka yang menarik dan dekat dengan kehidupan sehari-harimu untuk memantik rasa ingin tahu.
2. **Tahukah Kamu?:** Kotak fakta sains dan teknologi yang menambah wawasanmu tentang dunia komputer.
3. **Ayo Mencoba (Aktivitas Unplugged & Plugged):** Lembar kegiatan seru, baik bermain peran tanpa komputer maupun praktik langsung di depan laptop.
4. **Pojok Karakter & Etika:** Petuah bijak untuk membentuk pribadi yang jujur, santun, dan peduli sesama di dunia nyata maupun dunia digital.
5. **Rangkuman:** Ringkasan poin-poin penting yang sudah kalian pelajari di setiap bab.
6. **Uji Pemahaman (Evaluasi HOTS):** Latihan soal seru (pilihan ganda dan uraian) serta tugas proyek mini untuk menguji ketangkasan berpikirmu.
7. **Jurnal Refleksi:** Ruang untuk mencatat perasaan dan hal-hal baru yang sudah berhasil kalian kuasai.

---

## PETA KONSEP TAHUNAN KELAS 5 (FASE C)

```
                       PETUALANGAN KODING & KA KELAS V
                                      │
        ┌─────────────────────────────┴─────────────────────────────┐
        ▼                                                           ▼
   SEMESTER 1                                                  SEMESTER 2
┌───────────────────────────────┐           ┌───────────────────────────────┐
│ BAB 1: Berpikir Komputasional │           │ BAB 4: Mengenal Kecerdasan    │
│        dan Pemecahan Masalah  │           │        Artifisial (KA)        │
├───────────────────────────────┤           ├───────────────────────────────┤
│ BAB 2: Mengenal Teknologi     │           │ BAB 5: Eksperimen Klasifikasi │
│        Digital & Komputer     │           │        Data Bersama AI        │
├───────────────────────────────┤           ├───────────────────────────────┤
│ BAB 3: Menjaga Rahasia &      │           │ BAB 6: Berkreasi dengan Koding│
│        Keamanan Data Pribadi  │           │        Blok Visual & Gelar    │
└───────────────────────────────┘           │        Karya Siswa            │
                                            └───────────────────────────────┘
```

---

## DAFTAR TOKOH SAHABAT BELAJAR
* **Erik:** Murid kelas 5 yang penuh semangat, suka bertanya, dan gemar merakit mainan logika.
* **Siti:** Murid kelas 5 yang teliti, gemar merawat tanaman bunga, dan senang menggambar digital.
* **Made:** Murid kelas 5 yang lincah, pandai menari tarian daerah, dan kreatif mencari ide-ide baru.
* **Kibo si Robot Cerdas:** Karakter robot kartun berwarna biru-oranye yang ramah, sopan, dan setia mendampingi anak-anak belajar.
"""

save_md("Buku_Siswa_Koding_dan_KA_Kelas_5_SD_Part1.md", front_matter)
print("Part 1 saved.")
