# -*- coding: utf-8 -*-
import docx, re

def audit_file(filepath):
    print(f"\nAuditing: {filepath}")
    doc = docx.Document(filepath)
    stray_asterisks = 0
    stray_backticks = 0
    total_paragraphs = len(doc.paragraphs)
    total_tables = len(doc.tables)
    
    for i, p in enumerate(doc.paragraphs):
        t = p.text
        if "**" in t:
            stray_asterisks += 1
            print(f"  [P-{i}] Stray **: {t[:80]}")
        if "`" in t:
            stray_backticks += 1
            print(f"  [P-{i}] Stray backtick: {t[:80]}")
            
    for ti, tbl in enumerate(doc.tables):
        for ri, r in enumerate(tbl.rows):
            for ci, c in enumerate(r.cells):
                for pi, p in enumerate(c.paragraphs):
                    t = p.text
                    if "**" in t:
                        stray_asterisks += 1
                        print(f"  [T{ti} R{ri} C{ci}] Stray **: {t[:80]}")
                    if "`" in t:
                        stray_backticks += 1
                        print(f"  [T{ti} R{ri} C{ci}] Stray backtick: {t[:80]}")

    print(f"Result: Paragraphs={total_paragraphs}, Tables={total_tables}")
    print(f"Stray ** count: {stray_asterisks}")
    print(f"Stray ` count: {stray_backticks}")
    return stray_asterisks, stray_backticks

if __name__ == "__main__":
    f1 = r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-7-smp\Buku_Siswa_Koding_dan_KA_Kelas_7_SMP.docx"
    f2 = r"d:\Xampp2\htdocs\aktama-main\buku-ajar-kka\buku-kelas-7-smp\Buku_Panduan_Guru_Koding_dan_KA_Kelas_7_SMP.docx"
    a1, b1 = audit_file(f1)
    a2, b2 = audit_file(f2)
    if a1 == 0 and b1 == 0 and a2 == 0 and b2 == 0:
        print("\nALL SMP 7 CHECKS PERFECT! ZERO STRAY ASTERISKS & ZERO STRAY BACKTICKS.")
    else:
        print("\nSOME STRAY ARTIFACTS FOUND!")
